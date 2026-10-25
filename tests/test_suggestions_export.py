import json
from pathlib import Path

import openpyxl
from bson import ObjectId

from app.processors.tabular import Cell, Table
from app.schemas.result import ResultDocument
from app.services.export import convert
from app.services.profiling import profile_table
from app.services.suggestions import suggest_rules


def analysis_result(columns: list[str], rows: list[list[Cell]]) -> ResultDocument:
    """표를 진단해 분석 결과 문서를 만든다.

    파라미터:
        columns: 열 이름
        rows: 행
    """
    profile = profile_table(Table("csv", "utf-8", None, columns, iter(rows)))
    return ResultDocument(
        field=ObjectId(),
        job_id=ObjectId(),
        result_type="analysis",
        version=1,
        quality=profile.quality,
        analysis=profile.analysis,
    )


def test_suggest_rules_from_quality_issues() -> None:
    rows: list[list[Cell]] = [[" 홍길동", "30", "포항", f"자유 서술 {i}"] for i in range(3)]
    rows += [["김철수", "", "", ""], ["김철수", "", "", ""]]
    result = analysis_result(["name", "age", "city", "memo"], rows)

    suggested = [(s.rule.model_dump(exclude_none=True), s.reason) for s in suggest_rules(result)]
    rules = [rule for rule, _ in suggested]

    assert {"type": "trim_whitespace", "column": "name"} in rules
    assert {"type": "remove_duplicates"} in rules
    assert {"type": "fill_missing", "column": "age", "strategy": "median"} in rules
    assert {"type": "fill_missing", "column": "city", "strategy": "mode"} in rules
    assert all(reason for _, reason in suggested)


def test_no_mode_for_free_text_column() -> None:
    rows: list[list[Cell]] = [[f"서로 다른 민원 내용 {i}"] for i in range(60)]
    rows.append([""])
    result = analysis_result(["content"], rows)

    assert suggest_rules(result) == []


def test_convert_csv_to_json_and_xlsx(tmp_path: Path) -> None:
    source = tmp_path / "v1.csv"
    source.write_text("이름,나이\n홍길동,30\n김철수,\n", encoding="utf-8-sig")

    convert(source, tmp_path / "out.json", "json")
    assert json.loads((tmp_path / "out.json").read_text(encoding="utf-8")) == [
        {"이름": "홍길동", "나이": "30"},
        {"이름": "김철수", "나이": None},
    ]

    convert(source, tmp_path / "out.xlsx", "xlsx")
    sheet = openpyxl.load_workbook(tmp_path / "out.xlsx").worksheets[0]
    assert [list(r) for r in sheet.iter_rows(values_only=True)] == [
        ["이름", "나이"],
        ["홍길동", "30"],
        ["김철수", None],
    ]
