from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
import xlwt

from app.core.errors import AppError
from app.processors.encoding import detect_encoding
from app.processors.tabular import Cell, normalize_headers, open_table
from app.processors.text import open_text


def read_all(path: Path, extension: str) -> tuple[list[str], list[list[Cell]], int]:
    """표를 끝까지 읽어 (열, 행, 길이가 달랐던 행 수)를 돌려준다.

    파라미터:
        path: 파일 경로
        extension: 확장자
    """
    with open_table(path, extension) as table:
        rows = list(table.rows)
        return table.columns, rows, table.ragged_row_count


@pytest.mark.parametrize(
    ("encoding", "expected"),
    [("utf-8", "utf-8"), ("utf-8-sig", "utf-8-sig"), ("cp949", "cp949"), ("utf-16", "utf-16")],
)
def test_detect_encoding(tmp_path: Path, encoding: str, expected: str) -> None:
    path = tmp_path / "a.txt"
    path.write_bytes("민원 내용입니다\n가로등".encode(encoding))

    assert detect_encoding(path) == expected


def test_detect_encoding_rejects_unknown(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_bytes(b"ok\n" + b"\xff\xfe\xfd" * 3 + b"\x80")

    with pytest.raises(AppError) as info:
        detect_encoding(path)
    assert info.value.code == "ENCODING_ERROR"


def test_csv_cp949_with_quotes_and_newlines(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    content = '번호,내용,위도\n1,"가로등이\n꺼졌어요",36.01\n2,"쉼표, 포함",\n'
    path.write_bytes(content.encode("cp949"))

    columns, rows, ragged = read_all(path, "csv")

    assert columns == ["번호", "내용", "위도"]
    assert rows == [["1", "가로등이\n꺼졌어요", "36.01"], ["2", "쉼표, 포함", ""]]
    assert ragged == 0


@pytest.mark.parametrize(("extension", "delimiter"), [("csv", ";"), ("csv", "|"), ("tsv", "\t")])
def test_csv_delimiters(tmp_path: Path, extension: str, delimiter: str) -> None:
    path = tmp_path / f"a.{extension}"
    path.write_text(f"a{delimiter}b\n1{delimiter}2\n", encoding="utf-8")

    columns, rows, _ = read_all(path, extension)

    assert columns == ["a", "b"]
    assert rows == [["1", "2"]]


def test_csv_ragged_rows_are_fitted(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    path.write_text("a,b,c\n1,2\n1,2,3,4\n\n5,6,7\n", encoding="utf-8")

    columns, rows, ragged = read_all(path, "csv")

    assert rows == [["1", "2", None], ["1", "2", "3"], [], ["5", "6", "7"]]
    assert ragged == 2


def test_normalize_headers() -> None:
    columns, empty, duplicates = normalize_headers(["이름", "", None, "이름", "이름", " 나이 "])

    assert columns == ["이름", "column_2", "column_3", "이름_2", "이름_3", "나이"]
    assert empty == 2
    assert duplicates == ["이름", "이름"]


def test_empty_csv_has_no_header(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    path.write_text("", encoding="utf-8")

    with pytest.raises(AppError) as info:
        read_all(path, "csv")
    assert info.value.code == "EMPTY_TABLE"


def test_xlsx_first_sheet_with_types(tmp_path: Path) -> None:
    path = tmp_path / "a.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "민원"
    sheet.append(["접수일", "건수", "내용"])
    sheet.append([datetime(2026, 7, 14), 3, "가로등"])
    sheet.append([None, 2.5, None])
    workbook.create_sheet("두번째").append(["무시"])
    workbook.save(path)

    with open_table(path, "xlsx") as table:
        rows = list(table.rows)
        assert table.sheet == "민원"
    assert table.columns == ["접수일", "건수", "내용"]
    assert rows == [[datetime(2026, 7, 14), 3, "가로등"], [None, 2.5, None]]


@pytest.mark.parametrize("extension", ["xlsx", "xls"])
def test_excel_stored_without_extension(tmp_path: Path, extension: str) -> None:
    # 업로드 원본은 저장소에 확장자 없이(uploads/YYYY/MM/<fileId>) 저장된다
    source = tmp_path / f"a.{extension}"
    if extension == "xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["이름"])
        sheet.append(["홍길동"])
        workbook.save(source)
    else:
        book = xlwt.Workbook()
        book.add_sheet("s").write(0, 0, "이름")
        book.get_sheet(0).write(1, 0, "홍길동")
        book.save(str(source))
    stored = tmp_path / "6addfb33f4a5a49f7c5c1def"
    source.rename(stored)

    columns, rows, _ = read_all(stored, extension)

    assert columns == ["이름"]
    assert rows == [["홍길동"]]


def test_xls_first_sheet_with_types(tmp_path: Path) -> None:
    path = tmp_path / "a.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("민원")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for col, name in enumerate(["접수일", "건수", "비율", "내용"]):
        sheet.write(0, col, name)
    sheet.write(1, 0, datetime(2026, 7, 14), date_style)
    sheet.write(1, 1, 3)
    sheet.write(1, 2, 0.5)
    sheet.write(1, 3, "가로등")
    workbook.save(str(path))

    columns, rows, _ = read_all(path, "xls")

    assert columns == ["접수일", "건수", "비율", "내용"]
    assert rows == [[datetime(2026, 7, 14), 3, 0.5, "가로등"]]


def test_broken_xlsx_is_invalid(tmp_path: Path) -> None:
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"PK\x03\x04broken")

    with pytest.raises(AppError) as info:
        read_all(path, "xlsx")
    assert info.value.code == "INVALID_FILE"


@pytest.mark.parametrize(
    ("content", "columns", "rows"),
    [
        (
            '[{"a": 1, "b": "x"}, {"b": "y", "c": [1, 2]}]',
            ["a", "b", "c"],
            [[1, "x", None], [None, "y", "[1, 2]"]],
        ),
        (
            '{"이름": "홍길동", "주소": {"동": "수청동"}}',
            ["이름", "주소"],
            [["홍길동", '{"동": "수청동"}']],
        ),
        ("[1, 2, null]", ["value"], [[1], [2], [None]]),
    ],
)
def test_json_shapes(
    tmp_path: Path, content: str, columns: list[str], rows: list[list[Cell]]
) -> None:
    path = tmp_path / "a.json"
    path.write_text(content, encoding="utf-8")

    assert read_all(path, "json")[:2] == (columns, rows)


def test_open_text_strips_newlines(tmp_path: Path) -> None:
    path = tmp_path / "a.log"
    path.write_bytes("INFO 시작\r\n\r\nERROR 실패\r\n".encode("cp949"))

    with open_text(path, "log") as text:
        lines = list(text.lines)
        assert text.encoding == "cp949"
    assert lines == ["INFO 시작", "", "ERROR 실패"]
