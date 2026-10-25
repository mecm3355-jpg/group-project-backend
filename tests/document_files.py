"""테스트용 DOCX·PDF 파일을 만드는 도구."""

import io

import docx


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """문단과 (선택) 표 하나가 든 DOCX를 만든다. 표는 문단 뒤에 붙는다.

    파라미터:
        paragraphs: 문단 글자 목록(빈 문자열이면 빈 문단)
        table: 표의 행 목록(각 행은 칸 글자 목록)
    """
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(added.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    """페이지마다 ASCII 글자 줄을 쓴 최소 PDF를 만든다. 빈 목록인 페이지는 글자 층이 없다.

    파라미터:
        pages: 페이지별 글자 줄 목록
    """
    objects: list[bytes] = []
    page_ids = [3 + i * 2 for i in range(len(pages))]
    font_id = 3 + len(pages) * 2
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for page_id, lines in zip(page_ids, pages, strict=True):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {page_id + 1} 0 R >>".encode()
        )
        commands = "".join(
            f"BT /F1 12 Tf 72 {720 - 20 * n} Td ({line}) Tj ET\n" for n, line in enumerate(lines)
        ).encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(commands) + commands + b"endstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1))
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()
